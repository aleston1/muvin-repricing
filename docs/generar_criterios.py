from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment
import re

wb = Workbook()
H = Font(bold=True, color="FFFFFF"); HF = PatternFill("solid", fgColor="15181C")
SUB = PatternFill("solid", fgColor="EFECE4"); EST = PatternFill("solid", fgColor="FDF3D8"); FICHA = PatternFill("solid", fgColor="E5F4EA")
WR = Alignment(wrap_text=True, vertical="top")
def header(ws, cols, widths):
    ws.append(cols)
    for i, c in enumerate(ws[1], 1):
        c.font = H; c.fill = HF; c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[c.column_letter].width = widths[i-1]
    ws.row_dimensions[1].height = 34
    ws.freeze_panes = "B2" if len(cols) > 8 else "A2"

# ------------------------------------------------------------ 1. Cómo usar
ws = wb.active; ws.title = "Cómo usar"
ws.column_dimensions["A"].width = 120
txt = [
 ("Mapa de criterios del Asesor Muvin", True),
 ("", False),
 ("Esta planilla define CÓMO recomienda el asesor. El código la lee cada pocos minutos: lo que cambies acá se refleja en la web sin programar.", False),
 ("Precio y stock NO se cargan acá: salen en vivo de Tiendanube. Un producto solo se recomienda si tiene stock en la variante (talle/color) que le sirve al cliente.", False),
 ("", False),
 ("Pestañas", True),
 ("• Etapas: las preguntas que hace el asesor, en orden, con sus opciones y qué criterio alimenta cada una.", False),
 ("• Criterios: el diccionario de cada criterio (qué significa, si es filtro o suma puntaje, valores posibles).", False),
 ("• Productos: una fila por producto de Tiendanube (vehículos). Acá se marca bajo qué criterios se recomienda cada uno.", False),
 ("• Accesorios: qué accesorios se ofrecen en el combo según el tipo de vehículo y las respuestas.", False),
 ("• Talles: tabla de altura → talle de cuadro, y altura por rodado para infantiles.", False),
 ("• Listas: valores válidos para los desplegables.", False),
 ("", False),
 ("Colores en la pestaña Productos", True),
 ("Verde = dato tomado de la ficha del producto en Tiendanube.   Amarillo = estimado por tipo de producto: REVISAR.   Sin color = a completar.", False),
 ("", False),
 ("Reglas de oro", True),
 ("1. Si un producto no está en la pestaña Productos o tiene 'Activo en asesor' = No, el asesor nunca lo recomienda.", False),
 ("2. Productos nuevos: se agregan con su ID de Tiendanube (lo ves en la URL del admin). El asesor avisa en un log los productos con stock que no están mapeados.", False),
 ("3. 'Prioridad' (0 a 3) desempata: sirve para empujar lo que conviene vender (margen, sobrestock, lanzamientos).", False),
 ("4. Categorías nuevas (ej. patinetas, accesorios de camping): se agrega el tipo en Listas, sus preguntas en Etapas y sus productos en Productos.", False),
]
for t, b in txt:
    ws.append([t]); ws.cell(ws.max_row, 1).font = Font(bold=b, size=14 if b and ws.max_row == 1 else 11); ws.cell(ws.max_row, 1).alignment = WR

# ------------------------------------------------------------ 2. Etapas
ws = wb.create_sheet("Etapas")
header(ws, ["#", "Etapa", "Pregunta al cliente", "Opciones (lo que ve el cliente)", "Criterio que alimenta", "Cómo decide", "Se pregunta si…"],
       [4, 16, 34, 48, 22, 58, 30])
etapas = [
 (1, "Necesidad", "¿Qué querés resolver?",
  "Moverme todos los días (trabajo, estudio, trámites)\nPasear y disfrutar\nHacer deporte o salir a la naturaleza\nLlevar carga o a los chicos\nEs para un chico o chica",
  "Usos", "FILTRO: el producto tiene que tener marcado ese uso en Productos. Todavía no se habla de bici, monopatín ni moto.", "Siempre"),
 (2, "Distancia", "¿Cuántos km hacés por trayecto? (solo ida)",
  "Menos de 5 km\nEntre 5 y 15 km\nEntre 15 y 30 km\nMás de 30 km",
  "Autonomía (km) / Propulsión", "Eléctricos: FILTRO autonomía ≥ ida y vuelta × 1,5 (margen por frío, subidas, peso). Sin motor: más de 15 km suma puntaje a los eléctricos y a bicis de ruta/gravel.", "Necesidad ≠ chico"),
 (3, "Terreno", "¿Cómo es el recorrido?",
  "Llano, asfalto o bicisenda\nCon subidas o puentes\nTierra, ripio o senderos\nUn poco de todo",
  "Terreno apto / Subida máx (%)", "FILTRO: el terreno tiene que estar entre los aptos. Con subidas suma a eléctricos y a monopatines con subida máx ≥ 20%.", "Siempre"),
 (4, "Esfuerzo", "¿Cuánto querés pedalear?",
  "Quiero pedalear, también es ejercicio\nPedalear, pero con ayuda de un motor\nNada, quiero que me lleve",
  "Propulsión", "FILTRO: Pedal → bicis sin motor. Ayuda → bicis eléctricas (pedaleo asistido). Nada → monopatines, motos y e-bikes con acelerador. Es la pregunta que separa las categorías.", "Necesidad ≠ chico"),
 (5, "Velocidad y papeles", "¿Por dónde vas a circular?",
  "Bicisendas y calles tranquilas (hasta 25 km/h)\nCalles y avenidas (hasta 45 km/h)\nAvenidas rápidas y accesos (más de 45 km/h)",
  "Velocidad máx / Requiere licencia", "FILTRO por velocidad máx. Si elige más de 45 km/h se aclara que requiere licencia y patente, y se pregunta si la tiene o la sacaría.", "Esfuerzo = Nada"),
 (6, "Portabilidad", "¿Lo tenés que cargar o combinar con otro transporte?",
  "Lo subo por escaleras\nLo combino con tren, subte o auto\nLo guardo en planta baja, cochera o patio",
  "Plegable / Peso (kg)", "Escaleras: FILTRO peso ≤ 20 kg (suma si ≤ 15). Combinar: FILTRO plegable = Sí. Planta baja: sin restricción.", "Necesidad ≠ chico"),
 (7, "Estacionamiento", "¿Lo vas a dejar estacionado en la calle?",
  "Sí, en la calle\nNo, lo guardo adentro", "Accesorios (candado)", "Define el candado del combo (alta seguridad vs. liviano). No filtra vehículos.", "Siempre"),
 (8, "Noche", "¿Vas a circular de noche?",
  "Casi siempre de día\nSí, también de noche", "Accesorios (luces / reflectivos)", "Define luces y chaleco del combo. No filtra vehículos.", "Siempre"),
 (9, "Altura", "¿Cuánto medís?", "Deslizador 90–210 cm",
  "Altura mín / máx / Talles", "FILTRO: altura dentro del rango del producto. Si tiene talles, se busca una variante CON STOCK en el talle que corresponde (pestaña Talles).", "Siempre (para chicos: altura del chico)"),
 (10, "Peso", "¿Cuánto pesás aproximadamente?", "Menos de 70 kg\n70 a 90 kg\n90 a 110 kg\nMás de 110 kg",
  "Carga máx (kg)", "FILTRO: peso de la persona + 10 kg ≤ carga máxima. Clave en monopatines (algunos soportan solo 90 kg).", "Propulsión ≠ Pedal"),
 (11, "Presupuesto", "¿Cuánto querés invertir en total?", "Rangos + monto libre + 'No tengo tope'",
  "Precio (en vivo de TN)", "Se reserva lo mínimo de los accesorios esenciales y el resto es tope del vehículo. Hasta +10% se muestra avisando cuánto se pasa.", "Siempre"),
 (12, "Ya tengo", "¿Ya tenés alguno de estos?", "Casco / Candado / Luces (según tipo de vehículo)",
  "Accesorios", "Lo que ya tiene no entra en el combo.", "Siempre"),
 ("—", "Stock", "(automático)", "—", "Stock TN", "FILTRO final: solo productos publicados con stock en la variante que le sirve. Nada sin stock llega al cliente.", "Siempre"),
]
for e in etapas: ws.append(list(e))
for row in ws.iter_rows(min_row=2):
    for c in row: c.alignment = WR

# ------------------------------------------------------------ 3. Criterios
ws = wb.create_sheet("Criterios")
header(ws, ["Criterio (columna en Productos)", "Qué significa", "Tipo", "Valores", "Lo usa la etapa"], [26, 60, 14, 40, 18])
crit = [
 ("Tipo de vehículo", "Familia del producto. Define qué accesorios lleva el combo y cómo se presenta.", "Info", "Ver Listas", "—"),
 ("Propulsión", "Pedal = sin motor. Asistida = motor que ayuda al pedalear. Eléctrica = anda sola (acelerador).", "Filtro", "Pedal / Asistida / Asistida + acelerador / Eléctrica", "4"),
 ("Uso: diario / paseo / deporte / carga / niños", "Marcar Sí en cada necesidad para la que se recomienda.", "Filtro", "Sí / No", "1"),
 ("Terreno apto", "Dónde anda bien.", "Filtro", "Asfalto / Asfalto y mixto / Mixto y tierra / Todo terreno", "3"),
 ("Autonomía (km)", "Autonomía real a considerar (usar el valor de ficha).", "Filtro", "Número", "2"),
 ("Velocidad máx (km/h)", "Según ficha.", "Filtro", "Número", "5"),
 ("Requiere licencia/patente", "Si hace falta licencia de conducir y patentamiento.", "Filtro", "Sí / No / Revisar", "5"),
 ("Plegable", "Si se pliega para guardar o subir al transporte.", "Filtro", "Sí / No / Parcial", "6"),
 ("Peso (kg)", "Peso del vehículo. Clave para escaleras.", "Filtro", "Número", "6"),
 ("Carga máx (kg)", "Peso máximo soportado (persona + carga).", "Filtro", "Número", "10"),
 ("Altura mín / máx (cm)", "Rango de altura del usuario. Si tiene talles, el rango lo da la pestaña Talles.", "Filtro", "Número", "9"),
 ("Talles por variante", "Sí = el talle está en las variantes de TN y se usa la pestaña Talles.", "Filtro", "Sí / No", "9"),
 ("Subida máx (%)", "Pendiente máxima (monopatines y motos).", "Puntaje", "Número", "3"),
 ("Nivel", "Posicionamiento: ayuda a ordenar dentro del presupuesto.", "Puntaje", "Entrada / Medio / Alto", "11"),
 ("Prioridad", "Empujón comercial para desempatar (0 = neutro, 3 = máximo).", "Puntaje", "0 a 3", "—"),
 ("Activo en asesor", "No = nunca se recomienda (usados, repuestos, productos que no queremos empujar).", "Filtro", "Sí / No", "—"),
 ("Argumento de venta", "Frase corta que el asesor muestra en '¿Por qué te lo recomendamos?'. Si está vacía, se arma sola.", "Texto", "Texto libre", "—"),
]
for c in crit: ws.append(list(c))
for row in ws.iter_rows(min_row=2):
    for c in row: c.alignment = WR

# ------------------------------------------------------------ 4. Productos
ws = wb.create_sheet("Productos")
cols = ["ID Tiendanube", "Nombre", "Categoría en TN", "Tipo de vehículo", "Propulsión",
        "Uso: diario", "Uso: paseo", "Uso: deporte / naturaleza", "Uso: carga / chicos", "Uso: niños",
        "Terreno apto", "Autonomía (km)", "Velocidad máx (km/h)", "Requiere licencia/patente",
        "Plegable", "Peso (kg)", "Carga máx (kg)", "Altura mín (cm)", "Altura máx (cm)", "Talles por variante",
        "Subida máx (%)", "Nivel", "Prioridad", "Activo en asesor", "Argumento de venta", "Notas / revisar"]
header(ws, cols, [13, 30, 20, 20, 18, 9, 9, 11, 11, 9, 18, 11, 11, 12, 10, 9, 10, 10, 10, 10, 10, 10, 9, 10, 40, 44])

rows = []
def add(pid, nombre, cat, tipo, prop, usos, terreno, aut=None, vel=None, lic="No", pleg="No", peso=None, carga=None,
        amin=None, amax=None, talles="No", subida=None, nivel=None, activo="Sí", arg="", notas="", ficha=(), est=()):
    rows.append(dict(v=[pid, nombre, cat, tipo, prop] + [("Sí" if u in usos else "No") for u in ("diario","paseo","deporte","carga","ninos")] +
                    [terreno, aut, vel, lic, pleg, peso, carga, amin, amax, talles, subida, nivel, 0, activo, arg, notas],
                     ficha=set(ficha), est=set(est)))

# Monopatines (datos de ficha)
F = ("aut","vel","peso","carga","pleg","subida")
add(282264187,"Max G3","Monopatines","Monopatín eléctrico","Eléctrica",{"diario","paseo"},"Asfalto",80,45,"No","Sí",24.6,130,None,None,subida=30,nivel="Alto",ficha=F,arg="80 km de autonomía y doble suspensión hidráulica: el más completo para ir lejos todos los días.",notas="Velocidad máx 45 km/h: revisar normativa local de monopatines.")
add(282311179,"GT3","Monopatines","Monopatín eléctrico","Eléctrica",{"diario","paseo"},"Asfalto",95,50,"No","Sí",39.5,150,subida=30,nivel="Alto",ficha=F,notas="39,5 kg: no apto para subir escaleras. Velocidad 50 km/h: revisar normativa.")
add(297070699,"E2 PLUS E II","Monopatines","Monopatín eléctrico","Eléctrica",{"diario"},"Asfalto",25,25,"No","Sí",15,90,subida=12,nivel="Entrada",ficha=F,arg="Liviano (15 kg) y simple: ideal para trayectos cortos y para subirlo a casa.",notas="Carga máx 90 kg y subida 12%: solo llano.")
add(312062393,"X30","Monopatines","Monopatín eléctrico","Eléctrica",{"diario"},"Asfalto",50,25,"No","Sí",20,120,nivel="Medio",ficha=("aut","vel","peso","carga"),est=("pleg",),notas="Micro. Confirmar que es plegable.")
add(363243370,"F3 Pro","Monopatines","Monopatín eléctrico","Eléctrica",{"diario"},"Asfalto",70,32,"No","Sí",19.3,120,subida=24,nivel="Medio",ficha=F)
# Motos
M = ("aut","vel","lic","peso","carga")
add(260345249,"TC Max","Motos eléctricas","Moto eléctrica","Eléctrica",{"diario","paseo"},"Asfalto",110,100,"Sí","No",105,150,nivel="Alto",ficha=M,notas="Altura de asiento 77 cm. Envío a cotizar.")
add(260346171,"CUx","Motos eléctricas","Moto eléctrica (scooter)","Eléctrica",{"diario"},"Asfalto",65,60,"Sí","No",64,None,nivel="Medio",ficha=("aut","vel","lic","peso"),notas="Peso sin batería. Envío a cotizar.")
add(260347602,"F01","Motos eléctricas","Moto eléctrica (scooter)","Eléctrica",{"diario"},"Asfalto",80,77,"Sí","No",92,None,nivel="Medio",ficha=("aut","vel","lic","peso"),notas="Peso sin batería (+18 kg batería).")
add(260683064,"CPX","Motos eléctricas","Moto eléctrica (scooter)","Eléctrica",{"diario","carga"},"Asfalto",70,90,"Sí","No",None,None,nivel="Alto",ficha=("aut","vel","lic"),notas="Equivale a 125 cc. Ampliable a 2 baterías.")
add(263947898,"CPX (2 Baterías 74V)","Motos eléctricas","Moto eléctrica (scooter)","Eléctrica",{"diario","carga"},"Asfalto",140,90,"Sí","No",None,None,nivel="Alto",ficha=("aut","vel","lic"))
add(327905251,"CPX Explorer (2 baterías)","Motos eléctricas","Moto eléctrica (scooter)","Eléctrica",{"diario","carga"},"Asfalto",None,105,"Sí","No",None,None,nivel="Alto",ficha=("vel","lic"),notas="Completar autonomía (la ficha no la dice).")
add(327910096,"CITI","Motos eléctricas","Moto eléctrica","Eléctrica",{"diario"},"Asfalto",107,80,"Sí","No",None,None,nivel="Medio",ficha=("aut","vel","lic"))
# E-bikes
E = "Bicicletas eléctricas"
add(260342361,"2Fold 20",E,"Bici eléctrica plegable","Asistida + acelerador",{"diario","paseo"},"Asfalto",40,25,"No","Sí",17.5,115,140,185,ficha=("aut","vel","pleg","peso","carga","amin","amax"),nivel="Medio",arg="La e-bike plegable más liviana: 17,5 kg, se guarda bajo el escritorio.")
add(260343629,"HSD P9",E,"Bici eléctrica de carga","Asistida",{"diario","carga"},"Asfalto y mixto",120,None,"No","Sí",25.7,170,150,195,ficha=("aut","pleg","peso","carga","amin","amax"),nivel="Alto",arg="Lleva a los chicos o 60 kg atrás y se guarda vertical ocupando poco lugar.",notas="Carga total 170 kg / usuario hasta 120 kg.")
add(260344367,"2Fold 20 Fat",E,"Bici eléctrica plegable","Asistida + acelerador",{"diario","paseo"},"Asfalto y mixto",40,32,"No","Sí",20,115,140,185,ficha=("aut","vel","pleg","peso","carga","amin","amax"),nivel="Medio")
add(260347235,"Brina2 X1000",E,"Bici eléctrica urbana/MTB","Asistida + acelerador",{"diario","paseo","deporte"},"Asfalto y mixto",70,38,"Revisar","No",22,None,ficha=("aut","vel","peso"),nivel="Alto",notas="1000 W y 38 km/h: revisar si requiere patente.")
add(263963220,"Quick Haul P9",E,"Bici eléctrica de carga","Asistida",{"diario","carga"},"Asfalto",106,None,"No","No",23,150,145,195,ficha=("aut","peso","carga","amin","amax"),nivel="Alto",notas="Autonomía 53–106 km según modo. 145 cm con caño de asiento corto.")
add(264034402,"Nbd S5I",E,"Bici eléctrica urbana","Asistida",{"diario","paseo"},"Asfalto",118,None,"No","Parcial",None,120,ficha=("aut","carga"),nivel="Alto",notas="Autonomía 51–118 km. Confirmar plegado (ficha: 'tiempo de plegado 5 s').")
add(273218901,"Slim",E,"Bici eléctrica plegable","Asistida + acelerador",{"diario"},"Asfalto",40,30,"No","Sí",18,None,ficha=("aut","vel","pleg","peso"),nivel="Medio")
add(273255228,"X350",E,"Bici eléctrica plegable (rodado ancho)","Asistida + acelerador",{"diario","paseo"},"Todo terreno",35,40,"Revisar","Sí",None,None,ficha=("aut","vel","pleg"),nivel="Medio",notas="Doble suspensión, 20x4.0.")
add(273267951,"Enduro Pro X",E,"Bici eléctrica plegable (rodado ancho)","Asistida + acelerador",{"diario","paseo"},"Todo terreno",50,40,"Revisar","Sí",None,None,ficha=("aut","vel","pleg"),nivel="Medio")
add(291838862,"Rider FT01 - 750W",E,"Bici eléctrica tipo moto","Asistida + acelerador",{"diario","paseo"},"Todo terreno",70,45,"Revisar","No",49,120,ficha=("aut","vel","peso","carga"),nivel="Medio",notas="49 kg: no apta para escaleras.")
add(314027238,"Segway Xyber",E,"Bici eléctrica tipo moto","Asistida + acelerador",{"diario","paseo"},"Todo terreno",90,55,"Revisar","No",63,180,ficha=("aut","vel","peso","carga"),nivel="Alto")
add(314180547,"Chopper FT02",E,"Bici eléctrica tipo moto","Asistida + acelerador",{"diario","paseo"},"Todo terreno",75,60,"Revisar","No",None,120,ficha=("aut","vel","carga"),nivel="Medio",notas="La ficha dice 'sin patente ni registro' con 60 km/h: validar legalmente.")
add(321841921,"Vecocraft Foldy-E Einhell",E,"Bici eléctrica plegable","Asistida",{"diario"},"Asfalto",30,25,"No","Sí",None,None,ficha=("aut","vel","pleg"),nivel="Entrada",arg="Usa baterías de herramientas Einhell: repuesto barato y fácil.")
add(338519904,"Rider FT01 - 1000W",E,"Bici eléctrica tipo moto","Asistida + acelerador",{"diario","paseo"},"Todo terreno",70,45,"Revisar","No",49,120,ficha=("aut","vel","peso","carga"),nivel="Medio")
add(339104476,"HTK 1000M",E,"Bici eléctrica urbana","Asistida",{"diario"},"Asfalto",None,None,"No","No",23,None,ficha=("peso",),nivel="Entrada",notas="Ficha incompleta: completar autonomía y velocidad.")
add(347573632,"Ponyboy",E,"Bici eléctrica tipo moto","Asistida",{"diario","paseo"},"Asfalto y mixto",None,None,"Revisar","No",None,None,nivel="Medio",notas="Ficha incompleta: completar autonomía, velocidad, peso.")
add(349261766,"E-Krabi",E,"Bici eléctrica plegable","Asistida",{"diario","paseo"},"Asfalto",100,24,"No","Sí",21.4,120,160,200,ficha=("aut","vel","pleg","peso","carga","amin","amax"),nivel="Medio")
add(349261769,"E-GOA",E,"Bici eléctrica plegable","Asistida",{"diario"},"Asfalto",60,24,"No","Sí",19,100,150,190,ficha=("aut","vel","pleg","peso","carga","amin","amax"),nivel="Medio")
add(349261777,"E-Easy Fat",E,"Bici eléctrica plegable (rodado ancho)","Asistida",{"diario","paseo"},"Todo terreno",100,24,"No","Sí",21.8,120,160,200,ficha=("aut","vel","pleg","peso","carga","amin","amax"),nivel="Medio")
add(350943156,"Apine Trail E Shimano",E,"Bici eléctrica de montaña (eMTB)","Asistida",{"deporte"},"Mixto y tierra",70,25,"No","No",None,None,talles="Sí",ficha=("aut",),est=("vel",),nivel="Alto",notas="Autonomía 25–70 km. Nombre en TN con typo ('Apine').")
add(354117181,"Rift Zone E2 SHIMANO (2024)",E,"Bici eléctrica de montaña (eMTB)","Asistida",{"deporte"},"Mixto y tierra",None,25,"No","No",None,None,talles="Sí",est=("vel",),nivel="Alto",notas="Completar autonomía (batería 630 Wh).")
add(361026582,"City",E,"Bici eléctrica tipo moto","Asistida + acelerador",{"diario"},"Asfalto",65,45,"Revisar","No",None,120,ficha=("aut","vel","carga"),nivel="Medio")

# Bicis convencionales (estimado por tipo)
def conv(ids_nombres, cat, tipo, usos, terreno, pleg="No", talles="Sí", notas="", activo="Sí", amin=None, amax=None):
    for pid, n in ids_nombres:
        est = ("usos","terreno") + (("pleg",) if pleg != "No" else ()) + (("amin","amax") if amin else ())
        add(pid, n, cat, tipo, "Pedal", usos, terreno, pleg=pleg, talles=talles, amin=amin, amax=amax, activo=activo, notas=notas, est=est)

conv([(260343619,"Quick 6"),(260345186,"Street 700c"),(260346843,"Space"),(260348567,"Alight 2 DD Disc"),(283613255,"Kentfield 1"),
      (284230738,"Fairfax 1 (2025)"),(284243181,"Presidio 1 (2025)"),(284246048,"Muirwoods (2022)"),(284254494,"Kentfield 2"),
      (297493608,"Stinson 1"),(297519717,"Stinson 2"),(349261735,"Roadkiller Lady Disk"),(349261744,"I Am Single"),
      (349261748,"Flower Bud Fairy"),(349261756,"VHS Player"),(349261782,"Lucky Clover"),(349261785,"Lucky Clover Low"),
      (349261838,"Roadkiller Disk"),(350940557,"Kentfield 3"),(350940977,"Kentfield 1 ST (caño bajo)"),
      (350941779,"Fairfax 1 ST (caño bajo)"),(353200327,"Fairfax 3 (2025)"),(354117226,"Stinson 2 ST")],
     "Bicicletas > Urbanas", "Bici urbana", {"diario","paseo"}, "Asfalto")
add(349261743,"The Tentacle","Bicicletas > Urbanas + Dirt/Stunt","Bici dirt/stunt","Pedal",{"deporte"},"Todo terreno",talles="Sí",est=("usos","terreno"),notas="Está en Urbanas y en Dirt/Stunt: revisar categoría.")
conv([(260340429,"Link D8"),(260340659,"Node D8"),(260341040,"Link C8"),(260341051,"Node D7i"),(260341563,"Eclipse D16"),
      (260342182,"Link A7"),(349261704,"Hopper Mini"),(349261708,"Hopper XL"),(349261749,"Krabi V-brake"),(349261759,"Krabi Disk"),
      (349261760,"Lentus Mini"),(349261768,"GOA V-brake"),(349261771,"Easy Disk"),(349261778,"Easy Fat"),(349261789,"Easy 8"),
      (349261794,"Speed Disk"),(349261833,"Easy Fat Nexus"),(349261834,"Lentus 8")],
     "Bicicletas > Plegables", "Bici plegable", {"diario","paseo"}, "Asfalto", pleg="Sí", talles="No", amin=150, amax=190,
     notas="Completar peso (clave para escaleras) y rango de altura de ficha.")
conv([(283042882,"Nicasio (2023)"),(283050150,"Nicasio 1 (2025)"),(283064870,"Four Corners 1 Sword (2025)"),(297455735,"DSX 1 (2026)"),
      (297486793,"DSX 2 (2026)"),(345977305,"Nicasio+ Sword (650B)"),(349261732,"Wanderer"),(349261740,"The Lightning"),
      (350939845,"Nicasio 1 ST (caño bajo)"),(350940160,"Four Corners 1 Cues (2026)"),(359825099,"Esker 5.0")],
     "Bicicletas > Gravel", "Bici gravel", {"diario","paseo","deporte"}, "Asfalto y mixto")
add(351124297,"Gestalt (2026)","Bicicletas > Gravel + Ruta","Bici gravel/ruta","Pedal",{"deporte","diario"},"Asfalto y mixto",talles="Sí",est=("usos","terreno"))
add(349261840,"Mom's Favorite","Bicicletas > Gravel + MTB","Bici gravel/MTB","Pedal",{"deporte","paseo"},"Mixto y tierra",talles="Sí",est=("usos","terreno"),notas="Está en Gravel y MTB: revisar.")
conv([(260343873,"Trail 6 MTB 2021"),(260346819,"Cheetah"),(284282046,"Alpine Trail Carbon 1 (2024)"),(284296407,"El Roy (2023)"),
      (297298148,"Bolinas Ridge 1"),(297298255,"Bolinas Ridge 2"),(297318978,"Bobcat Trail 4 (2025)"),(297327810,"San Quentin 1"),
      (349261728,"Boys Don't Cry"),(349261738,"Lone Ranger"),(349261821,"Sunday"),(349261825,"Sunday Low"),(349261827,"Big Time"),
      (352623351,"Rift Zone 1 (2026)"),(353164248,"Rift Zone 2 29 (2025)"),(353164254,"Bobcat Trail 5 (2025)"),
      (354117068,"Pine Mountain 1 - 29"),(354117137,"Team Marin 1")],
     "Bicicletas > MTB", "Bici de montaña (MTB)", {"deporte"}, "Mixto y tierra")
conv([(284309508,"Alcatraz 1"),(354117074,"Alcatraz 2"),(354117085,"Alcatraz 24\" (2026)")],
     "Bicicletas > De Carga", "Bici dirt/stunt", {"deporte"}, "Todo terreno",
     notas="Están en 'De Carga' pero son bicis de dirt/salto: revisar categoría en TN.")
# Infantiles
def rango(n):
    m = re.search(r'(\d+(?:[.,]5)?)\s*"?(?: Race)?$', n.replace("\"",""))
    r = {"12":(85,100),"14":(95,110),"16":(105,120),"20":(115,135),"24":(130,150),"26":(145,165),"27.5":(150,175),"27,5":(150,175)}
    return r.get(m.group(1)) if m else None
for pid, n in [(260342484,"Power"),(260342490,"Jumper"),(304920554,"Bayview Trail 24\""),(305009753,"Bayview Trail 20\""),
               (349261711,"Bubble 14 Race"),(349261714,"Bubble 16 Race"),(349261718,"Bubble 27,5 Race"),(349261720,"Chloe 27.5 Race"),
               (349261721,"Bubble 24 Race"),(349261724,"Bubble 26 Race"),(349261795,"Bubble 20 Race"),(349261799,"Chloe 20 Race"),
               (349261804,"Chloe 24 Race"),(349261809,"Chloe 26 Race"),(349261814,"Chloe 16 Race"),(350552513,"Ant"),
               (353164252,"Rift Zone 26"),(354117262,"San Quentin 24\"")]:
    rg = rango(n)
    add(pid, n, "Vehículos para niños", "Bici infantil/juvenil", "Pedal", {"ninos"}, "Asfalto y mixto",
        amin=rg[0] if rg else None, amax=rg[1] if rg else None, est=("usos","amin","amax") if rg else ("usos",),
        notas="" if rg else "Completar rodado / rango de altura.")

KEY = {"aut":11,"vel":12,"lic":13,"pleg":14,"peso":15,"carga":16,"amin":17,"amax":18,"subida":20,"usos":(5,6,7,8,9),"terreno":10}
for r in rows:
    ws.append(r["v"])
    ri = ws.max_row
    for k in r["ficha"]:
        for ci in (KEY[k] if isinstance(KEY[k], tuple) else (KEY[k],)):
            ws.cell(ri, ci+1).fill = FICHA
    for k in r["est"]:
        for ci in (KEY[k] if isinstance(KEY[k], tuple) else (KEY[k],)):
            ws.cell(ri, ci+1).fill = EST
    ws.cell(ri, 25).alignment = WR; ws.cell(ri, 26).alignment = WR
ws.auto_filter.ref = f"A1:{ws.cell(1, len(cols)).column_letter}{ws.max_row}"
n = ws.max_row
def dv(formula, rng):
    d = DataValidation(type="list", formula1=formula, allow_blank=True); ws.add_data_validation(d); d.add(rng)
dv("=Listas!$A$2:$A$20", f"D2:D{n+200}")
dv("=Listas!$B$2:$B$6", f"E2:E{n+200}")
dv('"Sí,No"', f"F2:J{n+200}")
dv("=Listas!$C$2:$C$6", f"K2:K{n+200}")
dv('"Sí,No,Revisar"', f"N2:N{n+200}")
dv('"Sí,No,Parcial"', f"O2:O{n+200}")
dv('"Sí,No"', f"T2:T{n+200}")
dv('"Entrada,Medio,Alto"', f"V2:V{n+200}")
dv('"0,1,2,3"', f"W2:W{n+200}")
dv('"Sí,No"', f"X2:X{n+200}")
ws.cell(1,1).comment = Comment("ID del producto en Tiendanube (número en la URL del admin). Es la llave con la que el asesor cruza stock y precio en vivo.", "Asesor")

# ------------------------------------------------------------ 5. Accesorios
ws = wb.create_sheet("Accesorios")
header(ws, ["Tipo de vehículo", "Rol en el combo", "Categoría TN (ID)", "Combo", "Condición", "Nivel a elegir", "Por qué (texto al cliente)"],
       [26, 18, 28, 12, 30, 16, 70])
acc = [
 ("Bicis (todas)", "Casco", "Cascos > Para ciclismo (30169776)", "Esencial", "Si no tiene casco", "Medio", "Lo primero: protege tu cabeza en cada salida."),
 ("Bicis (todas)", "Candado", "Candados y accesorios (30169763), sin 'Específicos de moto'", "Esencial", "Estaciona en la calle", "Alto, sin cables", "En la calle, un U-lock, cadena o plegable de alta seguridad es lo que frena un robo."),
 ("Bicis (todas)", "Candado", "Candados y accesorios (30169763)", "Esencial", "Guarda adentro", "Económico", "Para paradas cortas alcanza con uno liviano."),
 ("Bicis sin motor", "Luces", "Luces > Para bicicleta (30169757)", "Esencial", "Siempre (mejor nivel si anda de noche)", "Medio / Alto de noche", "Para ver y, sobre todo, que te vean."),
 ("Bicis eléctricas", "Luces", "Luces > Para bicicleta (30169757)", "Completo", "Solo si la ficha no trae luces", "Medio", "Luz extra para ser más visible."),
 ("Bicis (todas)", "Inflador", "Infladores > De mano (38434548)", "Completo", "Siempre", "Medio", "Cubiertas bien infladas: menos pinchaduras."),
 ("Bicis urbanas / plegables", "Portapaquetes o canasto", "Portapaquetes (30169786) / Bolsos y canastos (33821153)", "Completo", "Uso diario o paseo", "Medio", "Para llevar tus cosas sin cargar la espalda."),
 ("Bicis gravel / MTB / ruta", "Caramañola + soporte", "Caramañolas (30169771)", "Completo", "Uso deporte o distancia > 15 km", "Medio", "Hidratación a mano."),
 ("Bicis infantiles", "Casco", "Cascos > Para ciclismo (30169776), solo infantiles", "Esencial", "Siempre", "Medio", "Para que aprenda con el casco puesto desde el primer día."),
 ("Monopatín eléctrico", "Casco", "Cascos > Para ciclismo (30169776)", "Esencial", "Si no tiene casco", "Medio", "A más de 20 km/h, el casco no es opcional."),
 ("Monopatín eléctrico", "Candado", "Candados y accesorios (30169763): cable o plegable", "Esencial", "Estaciona en la calle", "Medio", "Para atarlo cuando bajás a hacer algo."),
 ("Monopatín eléctrico", "Chaleco reflectivo", "Chalecos reflectivos (30170021)", "Esencial", "Circula de noche", "Económico", "El monopatín es bajo: que te vean los autos."),
 ("Monopatín eléctrico", "Porta celular", "Porta celulares > Para bicicleta (37360471)", "Completo", "Uso diario", "Medio", "Para usar el GPS sin sacar el celular del bolsillo."),
 ("Monopatín eléctrico", "Guantes", "Guantes > Para ciclismo (30169781)", "Completo", "Siempre", "Medio", "Más agarre y protección en las manos."),
 ("Moto eléctrica", "Casco de moto", "Cascos > Para moto (30170011)", "Esencial", "Siempre (obligatorio)", "Medio", "Obligatorio para circular."),
 ("Moto eléctrica", "Traba disco / candado de moto", "Candados > Específicos de moto (35512464)", "Esencial", "Estaciona en la calle", "Alto", "La traba de disco impide que se la lleven rodando."),
 ("Moto eléctrica", "Baúl", "Baúles (30170012)", "Completo", "Uso diario o carga", "Medio", "Para guardar casco, compras o la mochila."),
 ("Moto eléctrica", "Cubremanos", "Cubremanos (30170054)", "Completo", "Siempre", "Medio", "Manos secas y abrigadas en invierno."),
 ("Moto eléctrica", "Guantes de moto", "Guantes > Para moto (30247745)", "Completo", "Siempre", "Medio", "Protección y agarre."),
 ("Moto eléctrica", "Funda", "Fundas para moto (37279095)", "Completo", "Estaciona en la calle", "Medio", "Protege del sol y la lluvia."),
 ("Bici eléctrica tipo moto", "Casco", "Cascos > Para moto (30170011)", "Esencial", "Velocidad máx > 25 km/h", "Medio", "Anda rápido: te recomendamos casco de moto."),
]
for a in acc: ws.append(list(a))
for row in ws.iter_rows(min_row=2):
    for c in row: c.alignment = WR

# ------------------------------------------------------------ 6. Talles
ws = wb.create_sheet("Talles")
header(ws, ["Talle de cuadro", "Altura desde (cm)", "Altura hasta (cm)", "Equivalente MTB (pulgadas)", "Equivalente ruta/gravel (cm)"], [16, 16, 16, 24, 26])
for t in [("XXS",135,152,"—","—"),("XS",148,163,"13–14\"","< 50"),("S",158,172,"15–16\"","50–52"),("M",168,181,"17–18\"","53–55"),
          ("L",177,190,"19–20\"","56–58"),("XL",186,198,"21\"+","59+"),("XXL",194,210,"—","—")]: ws.append(list(t))
ws.append([]); ws.append(["Rodado (infantiles)", "Altura desde (cm)", "Altura hasta (cm)", "Edad aprox.", ""])
for c in ws[ws.max_row]: c.font = Font(bold=True)
for t in [('12"',85,100,"2–4 años"),('14"',95,110,"3–5 años"),('16"',105,120,"4–6 años"),('20"',115,135,"6–9 años"),
          ('24"',130,150,"8–12 años"),('26"',145,165,"11+ años"),('27,5"',150,175,"juvenil")]: ws.append(list(t))
ws.append([]); ws.append(["Los rangos se solapan a propósito: en el borde se ofrecen los dos talles. Ajustar por marca si hace falta."])

# ------------------------------------------------------------ 7. Listas
ws = wb.create_sheet("Listas")
header(ws, ["Tipo de vehículo", "Propulsión", "Terreno apto"], [36, 24, 20])
tipos = ["Bici urbana","Bici plegable","Bici gravel","Bici gravel/ruta","Bici gravel/MTB","Bici de ruta","Bici de montaña (MTB)","Bici dirt/stunt",
         "Bici de carga","Bici infantil/juvenil","Bici eléctrica urbana","Bici eléctrica plegable","Bici eléctrica plegable (rodado ancho)",
         "Bici eléctrica urbana/MTB","Bici eléctrica tipo moto","Bici eléctrica de carga","Bici eléctrica de montaña (eMTB)",
         "Monopatín eléctrico","Moto eléctrica","Moto eléctrica (scooter)"]
prop = ["Pedal","Asistida","Asistida + acelerador","Eléctrica"]
ter = ["Asfalto","Asfalto y mixto","Mixto y tierra","Todo terreno"]
for i in range(len(tipos)):
    ws.append([tipos[i], prop[i] if i < len(prop) else None, ter[i] if i < len(ter) else None])

wb.save(__import__("sys").argv[1])
print(len(rows), "productos")
